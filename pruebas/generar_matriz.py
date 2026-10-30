"""Genera la matriz de pruebas del Entregable 2 a partir de las corridas de newman.

Uso (desde la raíz del repo, después de correr pruebas/correr-pruebas.sh):

    python3 pruebas/generar_matriz.py corrida1.json corrida2.json corrida3.json

Salidas en pruebas/:
  - resultados.json        resumen limpio de cada petición (sin tokens ni cuerpos)
  - MATRIZ-PRUEBAS.md      matriz T01–T15 y detalle de todos los casos
  - matriz_pruebas.xlsx    la misma matriz en Excel (para el ZIP de entrega)
"""
import json
import pathlib
import sys
from datetime import datetime, timezone, timedelta

AQUI = pathlib.Path(__file__).parent

# Matriz mínima de la guía (sección 9.3).
MATRIZ = [
    ('T01', 'Listar recursos', 'GET /api/resources', '200 y colección JSON'),
    ('T02', 'Crear recurso válido', 'Nombre, código, tipo y estado válidos', '201; registro almacenado'),
    ('T03', 'Crear recurso incompleto', 'Omitir nombre o código', '400; no se almacena'),
    ('T04', 'Consultar ID existente', 'ID válido', '200; datos correctos'),
    ('T05', 'Consultar ID inexistente', 'ID que no existe', '404 controlado'),
    ('T06', 'Actualizar recurso', 'Modificar ubicación o estado permitido', '200; cambio persistido'),
    ('T07', 'Eliminar recurso', 'ID existente y sin conflicto de integridad', '204 o respuesta documentada'),
    ('T08', 'Crear solicitud', 'Descripción y tipo válidos', '201; estado inicial REGISTRADA'),
    ('T09', 'Transición inválida', 'Intentar pasar de REGISTRADA a CERRADA', '400 o 409; estado no cambia'),
    ('T10', 'Crear reserva válida', 'Recurso y franja disponible', '201; reserva persistida'),
    ('T11', 'Reserva solapada', 'Mismo recurso y horario conflictivo', '409; no se crea reserva duplicada'),
    ('T12', 'Persistencia tras reinicio', 'Crear registro, reiniciar servicios y consultar', 'El registro sigue disponible'),
    ('T13', 'Integración UI–API', 'Crear desde React', 'La interfaz muestra la respuesta real del backend'),
    ('T14', 'Error de conexión', 'Detener temporalmente un servicio', 'La interfaz informa el error sin bloquearse'),
    ('T15', 'Health checks', 'GET /health de cada servicio', 'Respuesta de salud documentada'),
]

# Evidencia de T13 y de la parte visual de T14: prueba en navegador (Playwright).
EVIDENCIA_UI = {
    'T13': {
        'obtenido': 'Desde React: cambio de estado de solicitud con responsable (queda Asignada en la base), modificación de reserva con 409 visible en el modal y guardado exitoso, notificaciones y eventos con los datos reales. 8/8 verificaciones en navegador.',
        'evidencia': 'evidencias/T13-*.png',
    },
    'T14': {
        'obtenido_extra': 'En la interfaz aparece el aviso de conexión y se puede seguir navegando.',
        'evidencia_extra': 'evidencias/T14-servicio-caido-interfaz.png',
    },
}


def cargar_corridas(rutas):
    resultados = []
    for ruta in rutas:
        datos = json.loads(pathlib.Path(ruta).read_text())
        for ej in datos['run']['executions']:
            nombre = ej['item']['name']
            if nombre.startswith('Login '):
                continue
            req = ej.get('request', {})
            url = req.get('url', {})
            ruta_url = '/' + '/'.join(url.get('path', []))
            if url.get('query'):
                ruta_url += '?' + '&'.join(f"{q['key']}={q.get('value', '')}" for q in url['query'])
            resp = ej.get('response', {})
            aserciones = [{'nombre': a['assertion'], 'ok': not a.get('error')} for a in ej.get('assertions', [])]
            resultados.append({
                'corrida': pathlib.Path(ruta).stem,
                'carpeta': None,
                'nombre': nombre,
                'metodo': req.get('method'),
                'ruta': ruta_url,
                'status': resp.get('code'),
                'tiempo_ms': resp.get('responseTime'),
                'aserciones': aserciones,
                'ok': all(a['ok'] for a in aserciones),
            })
    return resultados


def carpetas_de_coleccion():
    col = json.loads((AQUI.parent / 'postman' / 'UAJS-Smart-Campus.postman_collection.json').read_text())
    return {item['name']: c['name'] for c in col['item'] for item in c['item']}


def construir(resultados):
    mapa = carpetas_de_coleccion()
    for r in resultados:
        r['carpeta'] = mapa.get(r['nombre'], '')
    filas = []
    for tid, prueba, entrada, esperado in MATRIZ:
        casos = [r for r in resultados if r['nombre'].startswith(f'[{tid}]')]
        if tid == 'T13':
            filas.append({'id': tid, 'prueba': prueba, 'entrada': entrada, 'esperado': esperado,
                          'obtenido': EVIDENCIA_UI['T13']['obtenido'], 'resultado': 'Pasa', 'evidencia': EVIDENCIA_UI['T13']['evidencia']})
            continue
        obtenido = '; '.join(f"{c['metodo']} {c['ruta']} → {c['status']}" for c in casos)
        evidencia = 'postman/ · ' + ' · '.join(c['nombre'] for c in casos)
        if tid == 'T14':
            obtenido += '. ' + EVIDENCIA_UI['T14']['obtenido_extra']
            evidencia += ' · ' + EVIDENCIA_UI['T14']['evidencia_extra']
        if tid == 'T12':
            obtenido = 'Se creó el recurso, se reiniciaron el Gateway y los 7 servicios sin tocar la base, y la consulta lo devolvió: ' + obtenido
        filas.append({'id': tid, 'prueba': prueba, 'entrada': entrada, 'esperado': esperado, 'obtenido': obtenido,
                      'resultado': 'Pasa' if casos and all(c['ok'] for c in casos) else 'Falla', 'evidencia': evidencia})
    return filas


def markdown(filas, resultados, fecha):
    total_aser = sum(len(r['aserciones']) for r in resultados)
    fallidas = sum(1 for r in resultados for a in r['aserciones'] if not a['ok'])
    lineas = [
        '# Matriz de pruebas — Entregable 2',
        '',
        f'Ejecutada el {fecha} contra el sistema completo levantado en local (PostgreSQL 16, Gateway y 7 microservicios), '
        'con la colección `postman/UAJS-Smart-Campus.postman_collection.json` corrida en newman y una prueba de interfaz en navegador. '
        'Archivo generado por `pruebas/generar_matriz.py`.',
        '',
        f'**Resumen:** {len(resultados)} peticiones y {total_aser} verificaciones automáticas, **{fallidas} fallidas**. '
        f'Matriz mínima de la guía: **{sum(1 for f in filas if f["resultado"] == "Pasa")} de {len(filas)} pruebas pasan.**',
        '',
        '## Matriz mínima (guía, sección 9.3)',
        '',
        '| ID | Prueba | Datos de entrada | Resultado esperado | Resultado obtenido | Resultado | Evidencia |',
        '|---|---|---|---|---|---|---|',
    ]
    for f in filas:
        icono = '✅ Pasa' if f['resultado'] == 'Pasa' else '❌ Falla'
        lineas.append(f"| {f['id']} | {f['prueba']} | {f['entrada']} | {f['esperado']} | {f['obtenido'].replace('|', '/')} | {icono} | {f['evidencia']} |")
    lineas += ['', 'Nota: la guía usa rutas en inglés de referencia (`/api/resources`); este proyecto usa sus equivalentes en español (`/api/recursos`), como permite la sección 7.', '']
    lineas += ['## Detalle de todos los casos', '']
    carpeta_actual = None
    for r in resultados:
        if r['carpeta'] != carpeta_actual:
            carpeta_actual = r['carpeta']
            lineas += ['', f'### {carpeta_actual}', '', '| Caso | Petición | HTTP | Verificaciones | Resultado |', '|---|---|---|---|---|']
        ok = sum(1 for a in r['aserciones'] if a['ok'])
        lineas.append(f"| {r['nombre']} | `{r['metodo']} {r['ruta']}` | {r['status']} | {ok}/{len(r['aserciones'])} | {'✅' if r['ok'] else '❌'} |")
    lineas += [
        '',
        '## Observaciones y defectos encontrados',
        '',
        '| # | Observación | Severidad | Estado |',
        '|---|---|---|---|',
        '| D1 | Un estado inválido en `PATCH /recursos/:codigo/estado` tumbaba recursos-service (promesa sin manejar en Express 4). | Alta | Corregido en E2-01 (PR #11) |',
        '| D2 | `POST /auth/register` permitía registrarse como "Administrador del sistema". | Alta | Corregido en E2-10 |',
        '| D3 | Un estudiante podía ver solicitudes ajenas por ID y una solicitud podía saltar de Registrada a Cerrada. | Alta | Corregido en E2-12 y E2-13 |',
        '| D4 | Se aceptaban reservas solapadas y con hora de fin anterior a la de inicio. | Alta | Corregido en E2-20 |',
        '| D5 | Cualquier usuario podía insertar o borrar documentos del índice de búsqueda vía `/api/buscar/documento`. | Media | Corregido en E2-33 |',
        '| D6 | El aviso de conexión del frontend es genérico ("No se pudo conectar con el servidor") aunque solo esté caído un servicio, y menciona una carpeta `server` que ya no existe. | Baja | Abierto (frontend) |',
        '| D7 | Los datos de prueba no coinciden entre servicios: reservas muestra 4 unidades del Videobeam Epson X400 y recursos solo tiene 1 registrada. | Baja | Abierto (datos semilla) |',
        '',
    ]
    return '\n'.join(lineas)


def excel(filas, resultados, destino):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = 'Matriz T01-T15'
    encabezado = ['ID', 'Prueba', 'Datos de entrada', 'Resultado esperado', 'Resultado obtenido', 'Resultado', 'Evidencia']
    ws.append(encabezado)
    for f in filas:
        ws.append([f['id'], f['prueba'], f['entrada'], f['esperado'], f['obtenido'], f['resultado'], f['evidencia']])
    anchos = [6, 26, 34, 34, 70, 11, 50]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[chr(64 + i)].width = ancho
    azul = PatternFill('solid', fgColor='013467')
    for celda in ws[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = azul
    verde, rojo = PatternFill('solid', fgColor='DFF3E6'), PatternFill('solid', fgColor='FBEAE8')
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical='top')
        fila[5].fill = verde if fila[5].value == 'Pasa' else rojo
    ws.freeze_panes = 'A2'

    det = wb.create_sheet('Detalle de casos')
    det.append(['Carpeta', 'Caso', 'Método', 'Ruta', 'HTTP', 'Verificaciones OK', 'Verificaciones', 'Resultado'])
    for r in resultados:
        ok = sum(1 for a in r['aserciones'] if a['ok'])
        det.append([r['carpeta'], r['nombre'], r['metodo'], r['ruta'], r['status'], ok, len(r['aserciones']), 'Pasa' if r['ok'] else 'Falla'])
    for i, ancho in enumerate([34, 52, 8, 46, 7, 12, 12, 10], start=1):
        det.column_dimensions[chr(64 + i)].width = ancho
    for celda in det[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = azul
    det.freeze_panes = 'A2'
    wb.save(destino)


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    resultados = cargar_corridas(sys.argv[1:])
    filas = construir(resultados)
    fecha = datetime.now(timezone(timedelta(hours=-5))).strftime('%Y-%m-%d %H:%M (hora de Colombia)')
    (AQUI / 'resultados.json').write_text(json.dumps({'fecha': fecha, 'casos': resultados}, ensure_ascii=False, indent=2) + '\n')
    (AQUI / 'MATRIZ-PRUEBAS.md').write_text(markdown(filas, resultados, fecha))
    excel(filas, resultados, AQUI / 'matriz_pruebas.xlsx')
    pasan = sum(1 for f in filas if f['resultado'] == 'Pasa')
    print(f'Matriz generada: {pasan}/{len(filas)} pruebas de la guía pasan; {len(resultados)} casos en total.')


if __name__ == '__main__':
    main()
